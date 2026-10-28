"""Reads test data from the Excel workbook that drives the parametrized suites."""

from functools import lru_cache
from pathlib import Path

import pytest
from openpyxl import load_workbook

DATA_FILE = Path(__file__).parent / "saucedemo_test_data.xlsx"

SCHEMAS: dict[str, tuple[str, ...]] = {
    "users": ("id", "username", "password", "expected_result", "expected_message"),
    "products": ("id", "name", "price"),
    "sorting": ("id", "option", "expected_first_item", "expected_order"),
    "checkout": (
        "id",
        "first_name",
        "last_name",
        "postal_code",
        "expected_result",
        "expected_message",
    ),
    "http_checks": ("id", "path", "expected_status", "expected_content_type", "expected_text"),
}


@lru_cache(maxsize=None)
def load_rows(sheet: str) -> tuple[dict[str, str], ...]:
    """Return a sheet's data rows as dicts keyed by header, with blanks as ''."""
    if sheet not in SCHEMAS:
        raise KeyError(f"Unknown sheet '{sheet}'. Known sheets: {', '.join(SCHEMAS)}")
    if not DATA_FILE.is_file():
        raise FileNotFoundError(f"Test data workbook not found: {DATA_FILE}")

    workbook = load_workbook(DATA_FILE, read_only=True, data_only=True)
    try:
        raw_rows = [tuple(row) for row in workbook[sheet].iter_rows(values_only=True)]
    finally:
        workbook.close()

    headers = tuple("" if cell is None else str(cell).strip() for cell in raw_rows[0])
    missing = [column for column in SCHEMAS[sheet] if column not in headers]
    if missing:
        raise ValueError(f"Sheet '{sheet}' is missing columns: {', '.join(missing)}")

    return tuple(
        {header: "" if value is None else str(value).strip() for header, value in zip(headers, row)}
        for row in raw_rows[1:]
        if any(value is not None and str(value).strip() for value in row)
    )


def load_cases(sheet: str, *marks: str) -> list:
    """Return ``pytest.param`` objects, using each row's id as the test id."""
    return [
        pytest.param(row, id=row["id"], marks=[getattr(pytest.mark, mark) for mark in marks])
        for row in load_rows(sheet)
    ]
